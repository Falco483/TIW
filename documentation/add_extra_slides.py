import collections
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
import os

def add_server_side_slide(prs):
    slide_layout = prs.slide_layouts[1] # Titolo e Contenuto o 2 contenuti
    slide = prs.slides.add_slide(slide_layout)
    shapes = slide.shapes

    title_shape = shapes.title
    title_shape.text = "Server side: DAO & model objects"

    # Aggiungiamo 2 text box per simulare 2 colonne
    left_box = shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(5.5), Inches(5))
    tf_left = left_box.text_frame
    tf_left.word_wrap = True

    p = tf_left.paragraphs[0]
    p.text = "Controllers"
    p.font.size = Pt(28)
    p.level = 0

    controllers = [
        "LoginServlet",
        "ApiUserController",
        "ApiSkuController",
        "ApiProdottoTreeController",
        "ApiSyncController",
        "ApiConfigurazioneController",
        "ApiRicercaController",
        "ApiProdottoController"
    ]
    for c in controllers:
        p = tf_left.add_paragraph()
        p.text = c
        p.level = 1
        p.font.size = Pt(20)

    right_box = shapes.add_textbox(Inches(6.5), Inches(1.5), Inches(6), Inches(5))
    tf_right = right_box.text_frame
    tf_right.word_wrap = True

    p = tf_right.paragraphs[0]
    p.text = "Model objects (Beans)"
    p.font.size = Pt(22)
    p.level = 0

    beans = ["Utente", "SKU", "Prodotto", "Configurazione"]
    for b in beans:
        p = tf_right.add_paragraph()
        p.text = b
        p.level = 1
        p.font.size = Pt(18)

    p = tf_right.add_paragraph()
    p.text = "Data Access Objects (Classes)"
    p.font.size = Pt(22)
    p.level = 0

    daos = [
        ("UtenteDAO", ["checkCredentials(username, pwd)"]),
        ("SKUDAO", ["inserisciSku(...)", "getPrezzoReale(id)", "eliminaDefinitivamente(...)"]),
        ("ProdottoDAO", ["inserisciProdottoComposto(...)", "getAlberoProdottoByCodice(...)"]),
        ("ConfigurazioneDAO", ["inserisciTestata(...)", "inserisciDettagliBatch(...)"])
    ]

    for dao_name, methods in daos:
        p = tf_right.add_paragraph()
        p.text = dao_name
        p.level = 1
        p.font.size = Pt(18)
        for m in methods:
            p = tf_right.add_paragraph()
            p.text = m
            p.level = 2
            p.font.size = Pt(14)

def add_client_side_slide(prs):
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    shapes = slide.shapes

    title_shape = shapes.title
    title_shape.text = "Client side: view & view component"

    left_box = shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(6), Inches(5.5))
    tf_left = left_box.text_frame
    tf_left.word_wrap = True

    p = tf_left.paragraphs[0]
    p.text = "A differenza della progettazione con HTML puro, le viste vengono calcolate sul lato client"
    p.level = 0
    p.font.size = Pt(22)

    p = tf_left.add_paragraph()
    p.text = "Un componente client-side della vista"
    p.level = 0
    p.font.size = Pt(22)

    sub_bullets = [
        "mantiene eventuali dati del view model",
        "espone le funzioni che gestiscono gli eventi associati al componente e la funzione di utilità per associare gli eventi ai rispettivi gestori",
        "espone le funzioni che invocano i servizi del server in maniera asincrona ed eseguono il rendering dei dati provenienti dal server"
    ]
    for sb in sub_bullets:
        p = tf_left.add_paragraph()
        p.text = sb
        p.level = 1
        p.font.size = Pt(18)

    right_box = shapes.add_textbox(Inches(6.8), Inches(1.0), Inches(6.0), Inches(6.0))
    tf_right = right_box.text_frame
    tf_right.word_wrap = True

    p = tf_right.paragraphs[0]
    p.text = "Index"
    p.level = 0
    p.font.size = Pt(20)

    p = tf_right.add_paragraph()
    p.text = "Login form"
    p.level = 1
    p.font.size = Pt(16)

    p = tf_right.add_paragraph()
    p.text = "Gestione del submit e degli errori"
    p.level = 2
    p.font.size = Pt(14)

    components = [
        ("AppFornitore", [
            ("init:", "carica dati iniziali dal server (orfani, catalogo)"),
            ("handleSalvaTree:", "invia le azioni dell'editor albero al server"),
            ("handleSubmitSku:", "invia i dati del nuovo componente SKU")
        ]),
        ("AppCliente", [
            ("caricaCatalogo:", "richiede i prodotti al server e aggiorna la vista"),
            ("apriConfigurazione:", "renderizza l'albero per la scelta SKU"),
            ("handleSalvaConfigurazione:", "invia i dati della configurazione al server")
        ])
    ]

    for comp_name, methods in components:
        p = tf_right.add_paragraph()
        p.text = comp_name
        p.level = 0
        p.font.size = Pt(20)
        
        for m_name, m_desc in methods:
            p = tf_right.add_paragraph()
            p.text = f"{m_name} {m_desc}"
            p.level = 1
            p.font.size = Pt(14)

def main():
    filepath = r"c:\Users\Antonio\Desktop\progetto TIW\documentation\tabelle_tiw.pptx"
    prs = Presentation(filepath)
    add_server_side_slide(prs)
    add_client_side_slide(prs)
    prs.save(filepath)
    print("Slides added successfully to", filepath)

if __name__ == "__main__":
    main()
