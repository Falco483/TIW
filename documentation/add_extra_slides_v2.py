import os
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

def add_title_content_slide(prs, title, left_title, left_lines, right_title, right_lines):
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    shapes = slide.shapes

    title_shape = shapes.title
    title_shape.text = title

    # Left column
    left_box = shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(5.8), Inches(5.5))
    tf_left = left_box.text_frame
    tf_left.word_wrap = True

    if left_title:
        p = tf_left.paragraphs[0]
        p.text = left_title
        p.font.size = Pt(20)
        p.font.bold = True
        p.level = 0
        first = False
    else:
        first = True

    for text, level, size in left_lines:
        if first:
            p = tf_left.paragraphs[0]
            first = False
        else:
            p = tf_left.add_paragraph()
        p.text = text
        p.level = level
        p.font.size = Pt(size)

    # Right column
    right_box = shapes.add_textbox(Inches(6.5), Inches(1.5), Inches(6.5), Inches(5.5))
    tf_right = right_box.text_frame
    tf_right.word_wrap = True

    if right_title:
        p = tf_right.paragraphs[0]
        p.text = right_title
        p.font.size = Pt(20)
        p.font.bold = True
        p.level = 0
        first = False
    else:
        first = True

    for text, level, size in right_lines:
        if first:
            p = tf_right.paragraphs[0]
            first = False
        else:
            p = tf_right.add_paragraph()
        p.text = text
        p.level = level
        p.font.size = Pt(size)

def main():
    base_file = r"c:\Users\Antonio\Desktop\progetto TIW\documentation\tabelle_tiw.pptx"
    out_file = r"c:\Users\Antonio\Desktop\progetto TIW\documentation\tabelle_tiw_aggiornate.pptx"
    
    # Try to open the presentation or create a new one if it fails
    try:
        if os.path.exists(base_file):
            shutil.copy(base_file, out_file)
            prs = Presentation(out_file)
        else:
            prs = Presentation()
    except Exception as e:
        print(f"Cannot copy base file: {e}")
        prs = Presentation()

    # SLIDE 1: Server side: Controllers
    left_lines_1 = [
        ("API Controllers (SPA)", 0, 16),
        ("ApiConfigurazioneController", 1, 14),
        ("ApiProdottoController", 1, 14),
        ("ApiProdottoTreeController", 1, 14),
        ("ApiRicercaController", 1, 14),
        ("ApiSkuController", 1, 14),
        ("ApiSyncController", 1, 14),
        ("ApiUserController", 1, 14),
    ]
    right_lines_1 = [
        ("Servlets (SSR)", 0, 16),
        ("AzioneCatalogoServlet", 1, 14),
        ("CercaCatalogoServlet", 1, 14),
        ("ConfiguraServlet", 1, 14),
        ("ConfigurazioneActionServlet", 1, 14),
        ("DettaglioConfigurazioneServlet", 1, 14),
        ("HomeClienteServlet", 1, 14),
        ("HomeFornitoreServlet", 1, 14),
        ("LoginServlet", 1, 14),
        ("LogoutServlet", 1, 14),
        ("MieConfigurazioniServlet", 1, 14),
        ("RisultatoFornitoreServlet", 1, 14),
        ("SalvaConfigurazioneServlet", 1, 14)
    ]
    add_title_content_slide(prs, "Server side: Controllers", "SPA", left_lines_1, "SSR", right_lines_1)

    # SLIDE 2: Server side: DAO & Model objects
    left_lines_2 = [
        ("Beans", 0, 16),
        ("Configurazione", 1, 14),
        ("ElementoCatalogo", 1, 14),
        ("Prodotto", 1, 14),
        ("ProdottoComposto", 1, 14),
        ("ProdottoSemplice", 1, 14),
        ("SKU", 1, 14),
        ("Utente", 1, 14),
        ("DTOs", 0, 16),
        ("DettaglioDTO", 1, 14),
        ("UtenteSessionDTO", 1, 14),
        ("VoceConfigurazioneDTO", 1, 14),
        ("UtenteDAO", 0, 16),
        ("checkCredentials(username, password)", 1, 14),
        ("ConfigurazioneDAO", 0, 16),
        ("inserisciTestata(conf)", 1, 14),
        ("inserisciDettagliBatch(id, dettagli)", 1, 14),
        ("eliminaConfigurazione(id, user)", 1, 14),
        ("getConfigurazioneById(id, user)", 1, 14),
        ("getConfigurazioniByUtente(user)", 1, 14),
        ("updateTestata(conf)", 1, 14),
        ("eliminaConfigurazioniPerComponente(id, tipo)", 1, 14)
    ]
    right_lines_2 = [
        ("SKUDAO", 0, 16),
        ("findAll(), findById(id)", 1, 14),
        ("findByCodice(codice)", 1, 14),
        ("insert(sku), update(sku)", 1, 14),
        ("getPrezzoReale(id)", 1, 14),
        ("search(query)", 1, 14),
        ("eliminaDefinitivamente(id)", 1, 14),
        ("ProdottoDAO", 0, 16),
        ("getProdottiRadice()", 1, 14),
        ("findAll(), findById(id), findByCodice(cod)", 1, 14),
        ("getAlberoProdotto(id)", 1, 14),
        ("insertSemplice(...), insertComposto(...)", 1, 14),
        ("addSku(idProdotto, idSku)", 1, 14),
        ("addFiglio(idPadre, idFiglio)", 1, 14),
        ("search(query)", 1, 14),
        ("calcolaPrezziDaSku(id), ricalcolaPrezziComposto(id)", 1, 14),
        ("rimuoviAssociazioneSku(idProd, idSku)", 1, 14),
        ("rimuoviFiglio(idFiglio)", 1, 14),
        ("eliminaDefinitivamente(id)", 1, 14)
    ]
    add_title_content_slide(prs, "Server side: DAO & Model objects", "Model Objects", left_lines_2, "Data Access Objects", right_lines_2)

    # SLIDE 3: Client side: view & view component (Introduzione & Cliente)
    left_lines_3 = [
        ("A differenza della progettazione con HTML puro, le viste vengono calcolate sul lato client", 0, 16),
        ("Un componente client-side della vista:", 0, 16),
        ("mantiene eventuali dati del view model", 1, 14),
        ("espone le funzioni che gestiscono gli eventi associati al componente", 1, 14),
        ("espone le funzioni che invocano i servizi del server in maniera asincrona", 1, 14),
        ("esegue il rendering dei dati provenienti dal server", 1, 14),
        ("Index / Login form", 0, 16),
        ("Submit form (login.js) con fetch e redirect", 1, 14)
    ]
    right_lines_3 = [
        ("AppCliente", 0, 16),
        ("Inizializzazione & Navigazione", 1, 14),
        ("init, bindGlobalEvents, switchSection", 2, 12),
        ("Caricamento & Rendering Liste", 1, 14),
        ("caricaCatalogo, renderPaginaCatalogo", 2, 12),
        ("caricaConfigurazioni", 2, 12),
        ("Gestione Configurazione", 1, 14),
        ("apriConfigurazione, apriModifica, apriDettaglio", 2, 12),
        ("buildNodoConfigura, buildNodoModifica, buildNodoDettaglio", 2, 12),
        ("espandiNodo, raccogliScelte", 2, 12),
        ("Azioni Server (Async)", 1, 14),
        ("handleSalvaConfigurazione", 2, 12),
        ("cloneConfigurazione", 2, 12),
        ("eliminaConfigurazione", 2, 12),
        ("Utilità", 1, 14),
        ("mostraMessaggio, confermaAzione", 2, 12)
    ]
    add_title_content_slide(prs, "Client side: view & view component (1/2)", "Introduzione", left_lines_3, "Componente Cliente", right_lines_3)

    # SLIDE 4: Client side: view & view component (Fornitore)
    left_lines_4 = [
        ("Inizializzazione & Navigazione", 1, 14),
        ("init, bindGlobalEvents, switchSection", 2, 12),
        ("Creazione Componenti", 1, 14),
        ("handleSubmitSemplice, handleCreaComposto", 2, 12),
        ("handleSubmitSku, renderSkuCreata", 2, 12),
        ("Ricerca & Liste", 1, 14),
        ("handleSearch, handleExpandSearchResult", 2, 12),
        ("renderSkuCheckboxList, renderOrfaniCheckboxList", 2, 12),
        ("Gestione Prezzi", 1, 14),
        ("ricalcolaRiepilogoPrezzi, handleRicalcolaPrezzo", 2, 12)
    ]
    right_lines_4 = [
        ("Tree Editor: Rendering & Costruzione", 1, 14),
        ("renderTreeView, buildNodeUI, buildSkuUI", 2, 12),
        ("raccogliSkuDaAlbero, getNodeLevel", 2, 12),
        ("Tree Editor: Code di Azioni Locali", 1, 14),
        ("enqueueAction (aggiunge l'azione in RAM)", 2, 12),
        ("handleTreeAddChild, submitTreeAddChild", 2, 12),
        ("handleTreeAddSku, submitTreeAddSku", 2, 12),
        ("handleTreeUnlink, handleTreeUnlinkSku", 2, 12),
        ("handleTreeDelete, handleTreeDeleteSku", 2, 12),
        ("handleTreeInlineEdit, handleSkuInlineEdit", 2, 12),
        ("Azioni Server (Async)", 1, 14),
        ("handleSalvaTree (invia il batch di azioni pendenti)", 2, 12),
        ("handleEliminaSku, deleteProdottoFromDB", 2, 12),
        ("Utilità", 1, 14),
        ("mostraMessaggio, confermaAzione", 2, 12)
    ]
    add_title_content_slide(prs, "Client side: view & view component (2/2)", "AppFornitore", left_lines_4, "AppFornitore (Cont.)", right_lines_4)

    prs.save(out_file)
    print("Fatto. Salvato in:", out_file)

if __name__ == "__main__":
    main()
